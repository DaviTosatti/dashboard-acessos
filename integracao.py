"""Leitura das planilhas e publicação transacional da base do dashboard.

Uso na rotina existente:
    from integracao import atualizar_dashboard
    atualizar_dashboard('/caminho/dashboard_acessos/configuracao.json')

O consolidado é uma fotografia completa das entradas, não uma carga incremental.
"""
from __future__ import annotations

import argparse
import csv
import fnmatch
import hashlib
import json
import math
import re
import sqlite3
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from itertools import zip_longest
from pathlib import Path
from typing import Iterable

from padronizacao import AgentePadronizador, ErroPadronizacao, texto_cnpj, texto_identificador

PADRAO_CONFIG = Path(__file__).with_name('configuracao.json')


class ErroImportacao(ValueError):
    """A nova carga foi recusada. Uma base anterior permanece disponível."""


def normalizar(valor) -> str:
    texto = unicodedata.normalize('NFKD', str(valor or ''))
    return ''.join(c for c in texto if not unicodedata.combining(c)).casefold()


def chave_coluna(nome: str) -> str:
    return re.sub(r'[\W_]+', ' ', normalizar(nome).strip()).strip()


@dataclass(frozen=True)
class Configuracao:
    arquivo: Path
    valores: dict
    pasta: Path
    banco: Path


def carregar_configuracao(caminho=PADRAO_CONFIG) -> Configuracao:
    arquivo = Path(caminho).resolve()
    with arquivo.open(encoding='utf-8') as entrada:
        valores = json.load(entrada)
    if not isinstance(valores, dict):
        raise ErroImportacao('A configuração deve ser um objeto JSON.')
    for nome, padrao in [('linha_cabecalho', 1), ('limite_registros', 200000), ('limite_colunas', 128)]:
        valor = valores.setdefault(nome, padrao)
        if type(valor) is not int or valor < 1:
            raise ErroImportacao(f'{nome} deve ser um inteiro positivo.')
    if valores['limite_colunas'] > 256:
        raise ErroImportacao('O limite de colunas pode ser, no máximo, 256.')
    for nome in ['colunas_ocultas', 'colunas_data', 'colunas_numero', 'colunas_iniciais']:
        lista = valores.setdefault(nome, [])
        if not isinstance(lista, list) or not all(isinstance(v, str) for v in lista):
            raise ErroImportacao(f'{nome} deve ser uma lista de nomes.')
    abas = valores.setdefault('abas', None)
    if abas is not None and (not isinstance(abas, list) or not abas or not all(isinstance(a, str) for a in abas)):
        raise ErroImportacao('abas deve ser null ou uma lista não vazia de nomes de abas.')
    aliases = valores.setdefault('renomear_colunas', {})
    if not isinstance(aliases, dict) or not all(isinstance(k, str) and isinstance(v, str) and v.strip() for k, v in aliases.items()):
        raise ErroImportacao('renomear_colunas deve conter nomes de origem e destino não vazios.')
    padroes = valores.setdefault('padroes', ['*.xlsx', '*.xlsm', '*.csv', '*.tsv'])
    if not isinstance(padroes, list) or not padroes or not all(isinstance(p, str) for p in padroes):
        raise ErroImportacao('padroes deve ser uma lista não vazia.')
    separador = valores.setdefault('separador_csv', None)
    if separador is not None and (not isinstance(separador, str) or len(separador) != 1):
        raise ErroImportacao('separador_csv deve ser null ou um único caractere.')
    def resolver(nome, padrao):
        caminho = Path(valores.get(nome, padrao))
        return caminho.resolve() if caminho.is_absolute() else (arquivo.parent / caminho).resolve()
    return Configuracao(arquivo, valores, resolver('pasta_planilhas', 'planilhas_aprovadas'), resolver('banco', 'dados/dashboard.sqlite3'))


def arquivos_da_pasta(config: Configuracao) -> list[Path]:
    return sorted((p for p in config.pasta.rglob('*') if p.is_file()
                   and not p.name.startswith('~$')
                   and any(fnmatch.fnmatch(p.name.lower(), pat.lower()) for pat in config.valores['padroes'])),
                  key=lambda p: str(p).casefold())


def data_iso(valor) -> str:
    if isinstance(valor, datetime):
        return valor.isoformat(timespec='seconds')
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor).strip()
    try:
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}', texto):
            return date.fromisoformat(texto).isoformat()
        if re.fullmatch(r'\d{2}/\d{2}/\d{4}', texto):
            return datetime.strptime(texto, '%d/%m/%Y').date().isoformat()
        return datetime.fromisoformat(texto).isoformat(timespec='seconds')
    except ValueError as exc:
        raise ErroImportacao(f'Data inválida: {texto!r}. Use DD/MM/AAAA ou ISO.') from exc


def numero_texto(valor) -> str:
    texto = str(valor).strip()
    if isinstance(valor, str) and ',' in texto:
        texto = texto.replace('.', '').replace(',', '.')
    try:
        numero = Decimal(texto)
        if not numero.is_finite():
            raise InvalidOperation
        return format(numero, 'f')
    except InvalidOperation as exc:
        raise ErroImportacao(f'Número inválido: {valor!r}.') from exc


def transformar_valor(valor, tipo_forcado=None, formato='General'):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None, None
    if tipo_forcado == 'cnpj':
        return texto_cnpj(valor, formato), 'texto'
    if tipo_forcado == 'identificador':
        return texto_identificador(valor, formato), 'texto'
    if tipo_forcado == 'texto':
        if isinstance(valor, (int, float)) and not isinstance(valor, bool):
            texto, _ = transformar_valor(valor, None, formato)
            return texto, 'texto'
        return str(valor).strip(), 'texto'
    if tipo_forcado == 'data' or isinstance(valor, (date, datetime)):
        return data_iso(valor), 'data'
    if tipo_forcado == 'numero':
        return numero_texto(valor), 'numero'
    if isinstance(valor, bool):
        return ('Sim' if valor else 'Não'), 'texto'
    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and not math.isfinite(valor):
            raise ErroImportacao('A planilha contém um número não finito.')
        # Um formato Excel composto só por zeros é usado frequentemente em IDs.
        if re.fullmatch(r'0{2,}', formato) and float(valor).is_integer():
            return str(int(valor)).zfill(len(formato)), 'texto'
        return numero_texto(valor), 'numero'
    return str(valor).strip(), 'texto'


def linhas_excel(arquivo: Path, config: Configuracao):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ErroImportacao('Instale as dependências: python -m pip install -r requirements.txt') from exc
    dados = load_workbook(arquivo, read_only=True, data_only=True, keep_links=False)
    formulas = None
    try:
        formulas = load_workbook(arquivo, read_only=True, data_only=False, keep_links=False)
        abas = config.valores['abas'] or dados.sheetnames
        for nome in abas:
            if nome not in dados.sheetnames:
                raise ErroImportacao(f'{arquivo.name}: aba {nome!r} não encontrada.')
            aba = dados[nome]
            aba_formula = formulas[nome]
            aba.reset_dimensions()
            aba_formula.reset_dimensions()
            inicio = config.valores['linha_cabecalho']
            for linha, (cells, original) in enumerate(zip_longest(
                    aba.iter_rows(min_row=inicio), aba_formula.iter_rows(min_row=inicio), fillvalue=()), inicio):
                values, formats = [], []
                for cell, formula in zip_longest(cells, original):
                    valor = cell.value if cell else None
                    if formula is not None and formula.data_type == 'f' and valor is None:
                        raise ErroImportacao(f'{arquivo.name}, {nome}, {formula.coordinate}: fórmula sem resultado salvo. Recalcule e salve no Excel.')
                    if cell is not None and cell.data_type == 'e':
                        raise ErroImportacao(f'{arquivo.name}, {nome}, {cell.coordinate}: erro do Excel {valor}.')
                    values.append(valor)
                    formats.append(cell.number_format if cell else 'General')
                yield nome, linha, values, formats
    finally:
        dados.close()
        if formulas is not None:
            formulas.close()


def linhas_csv(arquivo: Path, config: Configuracao):
    with arquivo.open('r', encoding=config.valores.get('codificacao_csv', 'utf-8-sig'), newline='') as entrada:
        separador = config.valores['separador_csv']
        if separador is None:
            amostra = entrada.read(8192)
            entrada.seek(0)
            if arquivo.suffix.lower() == '.tsv':
                separador = '\t'
            else:
                try:
                    separador = csv.Sniffer().sniff(amostra, delimiters=';,\t|').delimiter
                except csv.Error:
                    separador = ';'
        leitor = csv.reader(entrada, delimiter=separador, strict=True)
        for linha, values in enumerate(leitor, 1):
            if linha >= config.valores['linha_cabecalho']:
                yield 'CSV' if arquivo.suffix.lower() == '.csv' else 'TSV', linha, values, ['General'] * len(values)


def atualizar_dashboard(config_path=PADRAO_CONFIG, planilhas: Iterable[str | Path] | None = None,
                       *, permitir_vazio: bool = False, somente_analisar: bool = False) -> dict:
    """Substitui a base inteira após ler e validar TODAS as entradas.

    Se planilhas for fornecido, essa lista representa o consolidado completo.
    Se qualquer entrada falhar, a transação é revertida e a base antiga fica intacta.
    Execute somente depois de salvar/fechar as planilhas e concluir a aprovação.
    """
    config = carregar_configuracao(config_path)
    arquivos = arquivos_da_pasta(config) if planilhas is None else sorted({Path(p).resolve() for p in planilhas})
    if not arquivos and not permitir_vazio:
        raise ErroImportacao('Nenhuma planilha encontrada. A base anterior foi mantida. Verifique pasta_planilhas.')
    if not somente_analisar:
        config.banco.parent.mkdir(parents=True, exist_ok=True)
    try:
        agente = AgentePadronizador(config)
    except ErroPadronizacao as exc:
        raise ErroImportacao(str(exc)) from exc
    ocultas = {chave_coluna(agente.canonico(c)) for c in config.valores['colunas_ocultas']}
    datas = {chave_coluna(agente.canonico(c)) for c in config.valores['colunas_data']}
    numeros = {chave_coluna(agente.canonico(c)) for c in config.valores['colunas_numero']}
    if datas & numeros:
        raise ErroImportacao('Uma coluna não pode estar em colunas_data e colunas_numero ao mesmo tempo.')
    colunas, por_chave, tipos, fontes = [], {}, {}, []
    quantidade = 0
    conn = sqlite3.connect(':memory:' if somente_analisar else config.banco, timeout=60)
    try:
        conn.execute('PRAGMA journal_mode=WAL')
        conn.execute('BEGIN IMMEDIATE')
        conn.execute('CREATE TABLE IF NOT EXISTS metadados (chave TEXT PRIMARY KEY, valor TEXT NOT NULL)')
        conn.execute('DROP TABLE IF EXISTS carga_temporaria')
        conn.execute('CREATE TABLE carga_temporaria (payload TEXT, arquivo TEXT, aba TEXT, linha INTEGER, tabela TEXT, chave_loja TEXT, origens_campos TEXT)')
        for arquivo in arquivos:
            if arquivo.suffix.lower() not in {'.xlsx', '.xlsm', '.csv', '.tsv'}:
                raise ErroImportacao(f'Formato não suportado: {arquivo.name}. Converta .xls para .xlsx.')
            antes = arquivo.stat()
            origem = arquivo.relative_to(config.pasta).as_posix() if arquivo.is_relative_to(config.pasta) else arquivo.name
            leitor = linhas_excel if arquivo.suffix.lower() in {'.xlsx', '.xlsm'} else linhas_csv
            cabecalhos, por_aba = {}, {}
            for aba, linha, valores, formatos in leitor(arquivo, config):
                if aba not in cabecalhos:
                    if not any(v is not None and str(v).strip() for v in valores):
                        raise ErroImportacao(f'{origem}, {aba}: cabeçalho vazio na linha {linha}.')
                    fonte = origem + '::' + aba
                    nomes_originais = [str(v).strip() if v is not None else '' for v in valores]
                    agente.registrar_cabecalho(fonte, nomes_originais)
                    campos = []
                    for valor in valores:
                        nome_original = str(valor).strip() if valor is not None else ''
                        nome, _ = agente.resolver(nome_original, fonte)
                        chave = chave_coluna(nome)
                        if not chave:
                            campos.append(None)
                            continue
                        if chave in ocultas or chave_coluna(nome_original) in ocultas:
                            campos.append(False)
                            if (fonte, nome_original) in agente.mapeamentos:
                                agente.mapeamentos[(fonte, nome_original)]['oculta'] = True
                            continue
                        if chave not in por_chave:
                            ident = 'c_' + hashlib.sha256(chave.encode()).hexdigest()[:16]
                            por_chave[chave] = ident
                            colunas.append({'id': ident, 'nome': nome, 'tipo': 'texto', 'formato': agente.tipos.get(nome)})
                            tipos[ident] = set()
                        tipo = agente.tipos.get(nome)
                        tipo = tipo or ('data' if chave in datas else 'numero' if chave in numeros else None)
                        campos.append((por_chave[chave], tipo, nome_original, nome))
                    if len(colunas) > config.valores['limite_colunas']:
                        raise ErroImportacao('A carga excede limite_colunas. Configure quais campos importar.')
                    cabecalhos[aba] = campos
                    por_aba[aba] = 0
                    continue
                if not any(v is not None and str(v).strip() for v in valores):
                    continue
                campos = cabecalhos[aba]
                payload = {}
                por_nome = {}
                fonte = origem + '::' + aba
                for i in range(max(len(valores), len(campos))):
                    valor = valores[i] if i < len(valores) else None
                    campo = campos[i] if i < len(campos) else None
                    if campo is None:
                        if valor is not None and str(valor).strip():
                            raise ErroImportacao(f'{origem}, {aba}, linha {linha}: valor em coluna sem cabeçalho.')
                        continue
                    if campo is False:
                        continue
                    ident, tipo_forcado, nome_original, nome = campo
                    try:
                        texto, tipo = transformar_valor(valor, tipo_forcado, formatos[i] if i < len(formatos) else 'General')
                    except (ErroImportacao, ErroPadronizacao) as exc:
                        agente.erro('VALOR_INVALIDO', f'{fonte}, linha {linha}: {nome} contém um valor incompatível com {tipo_forcado or "o tipo do campo"}.',
                                    fonte=fonte, linha=linha, coluna=nome)
                        texto, tipo = None, None
                    if ident in payload and texto is not None and payload[ident] is not None and payload[ident] != texto:
                        agente.erro('COLUNAS_DIVERGENTES', f'{fonte}, linha {linha}: colunas que apontam para {nome} têm valores diferentes.',
                                    fonte=fonte, linha=linha, coluna=nome)
                    else:
                        if texto is not None or ident not in payload:
                            payload[ident] = texto
                    por_nome[nome] = payload.get(ident)
                    agente.observar_campo(fonte, nome_original, texto, tipo)
                    if tipo is not None:
                        tipos[ident].add(tipo)
                if not any(v is not None for v in payload.values()):
                    continue
                quantidade += 1
                if quantidade > config.valores['limite_registros']:
                    raise ErroImportacao('A carga excede limite_registros. Nenhuma linha foi publicada.')
                agente.observar_linha(fonte, linha, por_nome)
                conn.execute('INSERT INTO carga_temporaria VALUES (?, ?, ?, ?, ?, ?, ?)',
                             (json.dumps(payload, ensure_ascii=False), origem, aba, linha,
                              agente.regra_fonte(fonte)['tabela'], por_nome.get(agente.chave), '{}'))
                por_aba[aba] += 1
            depois = arquivo.stat()
            if (antes.st_size, antes.st_mtime_ns) != (depois.st_size, depois.st_mtime_ns):
                raise ErroImportacao(f'{origem}: o arquivo mudou durante a leitura. Salve e repita a importação.')
            fontes.append({'arquivo': origem, 'registros': sum(por_aba.values()),
                           'abas': [{'nome': a, 'registros': n} for a, n in por_aba.items()],
                           'modificado_em': datetime.fromtimestamp(antes.st_mtime, timezone.utc).isoformat()})
        if planilhas is None and arquivos_da_pasta(config) != arquivos:
            raise ErroImportacao('A lista de arquivos mudou durante a leitura. Repita a importação.')
        if quantidade == 0 and not permitir_vazio and not agente.erros:
            raise ErroImportacao('Nenhum registro válido. A base anterior foi mantida. Use --permitir-vazio para uma limpeza intencional.')
        if agente.catalogo['enriquecimento'] is not None:
            from relacionamentos import enriquecer_carga
            quantidade = enriquecer_carga(conn, agente, colunas, por_chave)
            if quantidade == 0 and not permitir_vazio and not agente.erros:
                agente.erro('TABELA_PRINCIPAL_VAZIA', 'O enriquecimento produziria uma base vazia. A publicação foi recusada.')
        analise = agente.finalizar()
        if somente_analisar:
            return {'total': quantidade, 'fontes': fontes, 'analise_padronizacao': analise}
        agente.exigir_valido()
        from prospeccao import publicar_prospeccao
        resumo_prospeccao = publicar_prospeccao(conn, config)
        for coluna in colunas:
            conjunto = tipos[coluna['id']]
            coluna['tipo'] = next(iter(conjunto)) if len(conjunto) == 1 else 'texto'
        campos_sql = ', '.join('"' + c['id'] + '" TEXT' for c in colunas)
        conn.execute('DROP TABLE IF EXISTS registros')
        conn.execute('CREATE TABLE registros (id INTEGER PRIMARY KEY' + (', ' + campos_sql if campos_sql else '') + ', arquivo TEXT, aba TEXT, linha INTEGER, busca TEXT, origens_campos TEXT)')
        ordem = [c['id'] for c in colunas]
        insert = 'INSERT INTO registros (' + ','.join('"' + c + '"' for c in ordem + ['arquivo', 'aba', 'linha', 'busca', 'origens_campos']) + ') VALUES (' + ','.join('?' for _ in range(len(ordem) + 5)) + ')'
        cursor = conn.execute('SELECT payload, arquivo, aba, linha, origens_campos FROM carga_temporaria')
        for payload, arquivo, aba, linha, origens in cursor:
            registro = json.loads(payload)
            values = [registro.get(c) for c in ordem]
            conn.execute(insert, values + [arquivo, aba, linha, '\x1f'.join(normalizar(v) for v in values), origens])
        papeis = {papel: por_chave.get(chave_coluna(agente.canonico(config.valores.get('coluna_' + papel, papel.title()))))
                  for papel in ['cliente', 'sistema', 'status']}
        for campo in set(v for v in papeis.values() if v):
            conn.execute(f'CREATE INDEX "idx_{campo}" ON registros ("{campo}")')
        conn.execute('CREATE INDEX idx_arquivo ON registros (arquivo)')
        resumo = {'colunas': colunas, 'papeis': papeis, 'fontes': fontes, 'total': quantidade,
                  'atualizado_em': datetime.now(timezone.utc).isoformat(), 'versao': uuid.uuid4().hex,
                  'analise_padronizacao': analise}
        if resumo_prospeccao is not None:
            resumo['prospeccao'] = resumo_prospeccao
        conn.execute('INSERT OR REPLACE INTO metadados VALUES (?, ?)', ('snapshot', json.dumps(resumo, ensure_ascii=False)))
        conn.execute('DROP TABLE carga_temporaria')
        conn.commit()
        return resumo
    except ErroPadronizacao as exc:
        conn.rollback()
        raise ErroImportacao(str(exc)) from exc
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def main():
    parser = argparse.ArgumentParser(description='Atualiza o dashboard a partir das planilhas aprovadas.')
    parser.add_argument('--config', default=str(PADRAO_CONFIG))
    parser.add_argument('--permitir-vazio', action='store_true', help='Autoriza substituir a base por uma carga vazia.')
    args = parser.parse_args()
    try:
        resumo = atualizar_dashboard(args.config, permitir_vazio=args.permitir_vazio)
    except Exception as exc:
        parser.exit(1, f'Importação recusada: {exc}\n')
    print(f"Dashboard atualizado: {resumo['total']} registros, {len(resumo['fontes'])} arquivos.")


if __name__ == '__main__':
    main()
